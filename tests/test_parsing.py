import io

import pytest
from docx import Document
from openpyxl import Workbook

from app.parsing import extract_text


def test_extract_text_txt():
    assert extract_text("memo.txt", "こんにちは".encode()) == "こんにちは"


def test_extract_text_unsupported_extension():
    with pytest.raises(ValueError, match="対応していないファイル形式です"):
        extract_text("resume.pdf", b"dummy")


@pytest.mark.parametrize("filename", ["skill.xlsx", "skill.docx"])
def test_extract_text_corrupt_file_raises_value_error(filename):
    with pytest.raises(ValueError, match="読み込めませんでした"):
        extract_text(filename, b"not a real zip file")


def test_extract_text_xlsx():
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["名前", "スキル"])
    ws.append(["山田", "Python"])
    buf = io.BytesIO()
    wb.save(buf)

    text = extract_text("skill.xlsx", buf.getvalue())
    assert "# シート: Sheet1" in text
    assert "山田\tPython" in text


def test_extract_text_docx():
    doc = Document()
    doc.add_paragraph("経歴サマリ")
    buf = io.BytesIO()
    doc.save(buf)

    text = extract_text("skill.docx", buf.getvalue())
    assert "経歴サマリ" in text
